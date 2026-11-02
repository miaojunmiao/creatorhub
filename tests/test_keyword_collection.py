import asyncio
import json
import tempfile
import unittest
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook
from fastapi import HTTPException
from sqlmodel import select

from app import db
from app.browser.fetcher import (
    _douyin_search_item_matches,
    _sort_douyin_search_items,
    _douyin_search_needs_verification,
    douyin_search_empty_error,
    douyin_search_exception_error,
    decode_douyin_search_payloads,
    extract_search_awemes,
    fetch_douyin_search,
)
from app.config import Config
from app.engine.collection import (
    KeywordCollector,
    _sort_ks_search_items,
    _sort_xhs_search_items,
    _xhs_content_matches,
)
from app.platforms.douyin import Aweme, MediaItem
from app.main import (
    KeywordCollectionIn,
    _collection_content_dict,
    _collection_error_for_display,
    _collection_job_dict,
    _collection_keywords,
    export_keyword_collection,
    list_keyword_collection_comments,
    list_keyword_collections,
    create_keyword_collection,
    keyword_collection_content_media,
    keyword_collection_local_media,
    update_keyword_collection,
)
from app.models import (
    DouyinAccount,
    KeywordCollectionComment,
    KeywordCollectionContent,
    KeywordCollectionJob,
)
from app.reporting import build_keyword_collection_report


class SearchExtractionTests(unittest.TestCase):
    def test_douyin_stream_chunks_decode_as_independent_json_payloads(self):
        first = json.dumps({
            "data": [{"aweme_info": {
                "aweme_id": "stream-1",
                "video": {"play_addr": {}},
            }}],
        }, separators=(",", ":")).encode()
        ack = json.dumps({"ack": 1}, separators=(",", ":")).encode()
        raw = (f"{len(first):x}\r\n".encode() + first + b"\r\n"
               + f"{len(ack):x}\r\n".encode() + ack + b"\r\n0\r\n\r\n")

        payloads = decode_douyin_search_payloads(raw)

        self.assertEqual(len(payloads), 2)
        self.assertEqual(extract_search_awemes(payloads[0])[0]["aweme_id"],
                         "stream-1")

    def test_invalid_douyin_stream_framing_is_ignored(self):
        self.assertEqual(decode_douyin_search_payloads(b"zz\r\n{}\r\n"), [])

    def test_douyin_verify_check_payload_is_detected(self):
        self.assertTrue(_douyin_search_needs_verification({
            "data": [],
            "search_nil_info": {
                "search_nil_type": "verify_check",
                "search_nil_item": "verify_check",
            },
        }))
        self.assertFalse(_douyin_search_needs_verification({
            "data": [{"aweme_info": {"aweme_id": "1"}}],
        }))

    def test_collection_error_display_shortens_closed_browser_trace(self):
        message = _collection_error_for_display(
            "ai短剧: 打开抖音搜索页失败: TargetClosedError('Target page, "
            "context or browser has been closed')")
        self.assertEqual(
            message,
            "ai短剧：采集窗口已关闭，请点击“续跑”并保持窗口开启",
        )
        self.assertNotIn("TargetClosedError", message)

    def test_captcha_page_has_actionable_error(self):
        message = douyin_search_empty_error(
            "验证码中间页", "", "https://www.douyin.com/search/example", [])
        self.assertIn("打开浏览器", message)
        self.assertIn("续跑", message)

    def test_login_modal_is_not_misclassified_as_security_captcha(self):
        message = douyin_search_empty_error(
            "抖音搜索", "扫码登录 验证码登录 登录后即可搜索更多精彩视频",
            "https://www.douyin.com/search/example", [])
        self.assertEqual(message, "抖音登录态已失效；请重新扫码登录后再续跑")

    def test_closed_visible_window_has_short_actionable_error(self):
        message = douyin_search_exception_error(
            RuntimeError("TargetClosedError: Target page, context or browser has been closed"))
        self.assertEqual(
            message,
            "抖音采集窗口被关闭；请续跑任务，并在任务结束前保持窗口打开",
        )

    def test_keyword_input_is_split_normalized_and_deduplicated(self):
        self.assertEqual(
            _collection_keywords(["露营, 骑行", "露营\n穿搭", "骑行"]),
            ["露营", "骑行", "穿搭"],
        )

    def test_collection_progress_tracks_execution_instead_of_result_count(self):
        job = KeywordCollectionJob(
            keywords='["露营", "骑行"]',
            max_contents_per_keyword=20,
            status="running",
            processed_keyword_count=1,
            current_item_index=2,
            current_item_total=4,
            content_count=0,
        )

        payload = _collection_job_dict(job)

        self.assertEqual(payload["total_keyword_count"], 2)
        self.assertEqual(payload["processed_keyword_count"], 1)
        self.assertEqual(payload["progress_percent"], 75)
        self.assertEqual(payload["content_count"], 0)
        job.status = "done"
        self.assertEqual(_collection_job_dict(job)["progress_percent"], 100)

    def test_extracts_known_search_wrappers_and_deduplicates(self):
        first = {
            "aweme_id": "a-1",
            "desc": "first",
            "video": {"play_addr": {"url_list": ["https://media.test/a.mp4"]}},
        }
        duplicate = {**first, "desc": "duplicate"}
        second = {
            "aweme_id": "a-2",
            "images": [{"url_list": ["https://media.test/b.jpg"]}],
        }
        payload = {
            "data": [
                {"aweme_info": first},
                {"search_result": {"aweme_mix_info": duplicate}},
                {"card": {"items": [second]}},
                {"aweme_info": {"aweme_id": "user-card-without-media"}},
            ]
        }

        result = extract_search_awemes(payload)

        self.assertEqual([row["aweme_id"] for row in result], ["a-1", "a-2"])
        self.assertEqual(result[0]["desc"], "first")

    def test_search_filters_and_sorting_are_applied_deterministically(self):
        rows = [
            {"aweme_id": "old", "create_time": 100, "video": {"play_addr": {}},
             "statistics": {"digg_count": 100, "comment_count": 8}},
            {"aweme_id": "new", "create_time": 190, "images": [{}],
             "statistics": {"digg_count": 20, "comment_count": 3}},
            {"aweme_id": "hot", "create_time": 180, "video": {"play_addr": {}},
             "statistics": {"digg_count": 500, "comment_count": 20}},
        ]

        self.assertFalse(_douyin_search_item_matches(
            rows[0], content_type="video", publish_time="day",
            min_likes=50, min_comments=5, now=100_000,
        ))
        self.assertTrue(_douyin_search_item_matches(
            rows[2], content_type="video", min_likes=100, min_comments=10,
        ))
        self.assertEqual(
            [row["aweme_id"] for row in _sort_douyin_search_items(rows, "latest")],
            ["new", "hot", "old"],
        )
        self.assertEqual(
            [row["aweme_id"] for row in _sort_douyin_search_items(rows, "most_liked")],
            ["hot", "old", "new"],
        )


class DouyinPageSearchTests(unittest.IsolatedAsyncioTestCase):
    async def test_repeated_challenges_share_a_budget_and_keep_partial_result_error(self):
        response = SimpleNamespace(
            status=200,
            url="https://www.douyin.com/aweme/v1/web/general/search/single/",
        )
        verify = {"data": [], "search_nil_info": {"search_nil_type": "verify_check"}}
        response.json = AsyncMock(return_value=verify)
        page = SimpleNamespace(
            url="about:blank", on=MagicMock(),
            wait_for_load_state=AsyncMock(), bring_to_front=AsyncMock(),
            close=AsyncMock(), evaluate=AsyncMock(),
        )
        verification_waits = 0

        async def emit(payload):
            response.json.return_value = payload
            await page.on.call_args.args[1](response)

        async def goto(url, **_kwargs):
            page.url = url
            await emit(verify)

        async def wait(milliseconds):
            nonlocal verification_waits
            if milliseconds == 1000:
                verification_waits += 1
                if verification_waits == 1:
                    await emit({"data": [{"aweme_info": {
                        "aweme_id": "partial-result", "video": {"play_addr": {}},
                    }}]})

        async def wheel(*_args):
            await emit(verify)

        page.goto, page.wait_for_timeout = goto, wait
        page.mouse = SimpleNamespace(wheel=AsyncMock(side_effect=wheel))
        rows, error = await fetch_douyin_search(
            SimpleNamespace(), SimpleNamespace(), "露营", max_results=10,
            captcha_wait_seconds=3, context=SimpleNamespace(pages=[page]),
        )
        self.assertEqual([row["aweme_id"] for row in rows], ["partial-result"])
        self.assertIn("请完成验证", error)
        self.assertEqual(verification_waits, 3)
        page.mouse.wheel.assert_awaited_once()
        page.close.assert_awaited_once()

    async def test_empty_search_stops_at_stagnant_limit(self):
        page = SimpleNamespace(
            url="about:blank", on=MagicMock(), goto=AsyncMock(),
            wait_for_load_state=AsyncMock(), wait_for_timeout=AsyncMock(),
            close=AsyncMock(), evaluate=AsyncMock(),
            title=AsyncMock(return_value="抖音搜索"),
            locator=MagicMock(return_value=SimpleNamespace(inner_text=AsyncMock(return_value=""))),
            mouse=SimpleNamespace(wheel=AsyncMock()),
        )
        with patch("app.browser.fetcher._submit_douyin_search", AsyncMock(return_value=False)):
            rows, error = await fetch_douyin_search(
                SimpleNamespace(), SimpleNamespace(), "露营", max_scrolls=40,
                stagnant_limit=2, context=SimpleNamespace(pages=[page]),
            )
        self.assertEqual(rows, [])
        self.assertTrue(error)
        self.assertEqual(page.mouse.wheel.await_count, 2)
        page.wait_for_load_state.assert_not_awaited()

    async def test_cancel_during_verification_stops_without_scrolling(self):
        progress = []
        canceled = False
        response = SimpleNamespace(
            status=200,
            url="https://www.douyin.com/aweme/v1/web/general/search/single/",
            json=AsyncMock(return_value={
                "data": [], "search_nil_info": {"search_nil_type": "verify_check"},
            }),
        )
        page = SimpleNamespace(
            url="about:blank", on=MagicMock(),
            wait_for_load_state=AsyncMock(), bring_to_front=AsyncMock(),
            close=AsyncMock(), mouse=SimpleNamespace(wheel=AsyncMock()),
        )

        async def goto(url, **_kwargs):
            page.url = url
            await page.on.call_args.args[1](response)

        async def wait(milliseconds):
            nonlocal canceled
            if milliseconds == 1000:
                canceled = True

        page.goto = goto
        page.wait_for_timeout = wait
        rows, error = await fetch_douyin_search(
            SimpleNamespace(), SimpleNamespace(), "露营",
            context=SimpleNamespace(pages=[page]),
            on_progress=progress.append, cancel_requested=lambda: canceled,
        )
        self.assertEqual((rows, error), ([], ""))
        self.assertTrue(any("等待人工" in step for step in progress))
        page.close.assert_awaited_once()
        page.mouse.wheel.assert_not_awaited()

    async def test_verification_recovery_with_filtered_results_does_not_keep_waiting(self):
        progress = []
        response = SimpleNamespace(
            status=200,
            url="https://www.douyin.com/aweme/v1/web/general/search/single/",
        )
        response.json = AsyncMock(return_value={
            "data": [], "search_nil_info": {"search_nil_type": "verify_check"},
        })
        page = SimpleNamespace(
            url="about:blank", on=MagicMock(),
            wait_for_load_state=AsyncMock(), bring_to_front=AsyncMock(),
            close=AsyncMock(), mouse=SimpleNamespace(wheel=AsyncMock()),
            evaluate=AsyncMock(), title=AsyncMock(return_value="抖音搜索"),
            locator=MagicMock(return_value=SimpleNamespace(inner_text=AsyncMock(return_value=""))),
        )
        verification_waits = 0

        async def goto(url, **_kwargs):
            page.url = url
            await page.on.call_args.args[1](response)

        async def wait(milliseconds):
            nonlocal verification_waits
            if milliseconds == 1000:
                verification_waits += 1
                response.json.return_value = {"data": [{"aweme_info": {
                    "aweme_id": "filtered", "video": {"play_addr": {}},
                    "statistics": {"digg_count": 1},
                }}]}
                await page.on.call_args.args[1](response)

        page.goto = goto
        page.wait_for_timeout = wait
        with patch("app.browser.fetcher._submit_douyin_search", AsyncMock(return_value=False)):
            rows, error = await fetch_douyin_search(
                SimpleNamespace(), SimpleNamespace(), "露营", min_likes=100,
                captcha_wait_seconds=5, max_scrolls=2,
                context=SimpleNamespace(pages=[page]), on_progress=progress.append,
            )
        self.assertEqual(rows, [])
        self.assertIn("没有符合条件", error)
        self.assertNotIn("请完成验证", error)
        self.assertEqual(verification_waits, 1)

    async def test_stream_response_body_is_decoded_and_collected(self):
        payload = json.dumps({"data": [{"aweme_info": {
            "aweme_id": "stream-result-1",
            "desc": "流式首屏",
            "video": {"play_addr": {"url_list": ["https://media.test/1.mp4"]}},
        }}]}, separators=(",", ":")).encode()
        framed = (f"{len(payload):x}\r\n".encode() + payload
                  + b"\r\n0\r\n\r\n")

        class Response:
            status = 200
            url = "https://www.douyin.com/aweme/v1/web/general/search/stream/"

            async def json(self):
                raise ValueError("chunk framed")

            async def body(self):
                return framed

        class Page:
            def __init__(self):
                self.url = "about:blank"
                self.handler = None
                self.mouse = SimpleNamespace(wheel=AsyncMock())

            def on(self, event, handler):
                if event == "response":
                    self.handler = handler

            async def goto(self, url, **_kwargs):
                self.url = url
                await self.handler(Response())

            async def wait_for_load_state(self, *_args, **_kwargs):
                return None

            async def wait_for_timeout(self, _timeout):
                return None

            async def close(self):
                return None

        page = Page()
        rows, error = await fetch_douyin_search(
            SimpleNamespace(), SimpleNamespace(), "数码", max_results=1,
            context=SimpleNamespace(pages=[page]))

        self.assertEqual([row["aweme_id"] for row in rows],
                         ["stream-result-1"])
        self.assertEqual(error, "")
        page.mouse.wheel.assert_not_awaited()

    async def test_verification_wait_is_passive_and_page_response_resumes_collection(self):
        verify_payload = {
            "data": [],
            "search_nil_info": {"search_nil_type": "verify_check"},
        }
        success_payload = {"data": [{"aweme_info": {
            "aweme_id": "page-result-1",
            "desc": "页面自然响应",
            "video": {"play_addr": {"url_list": ["https://media.test/1.mp4"]}},
        }}]}

        class Response:
            status = 200
            url = "https://www.douyin.com/aweme/v1/web/general/search/single/"

            def __init__(self, payload):
                self.payload = payload

            async def json(self):
                return self.payload

        class Page:
            def __init__(self):
                self.url = "about:blank"
                self.handler = None
                self.wait_calls = 0
                self.evaluate_calls = 0
                self.mouse = SimpleNamespace(wheel=AsyncMock())

            def on(self, event, handler):
                if event == "response":
                    self.handler = handler

            async def goto(self, url, **_kwargs):
                self.url = url
                await self.handler(Response(verify_payload))

            async def wait_for_load_state(self, *_args, **_kwargs):
                return None

            async def wait_for_timeout(self, _timeout):
                self.wait_calls += 1
                # 第一次是页面 settle；第二次是验证码被动等待。模拟用户完成验证后，
                # 页面自己产生正常搜索响应。
                if self.wait_calls == 2:
                    await self.handler(Response(success_payload))

            async def bring_to_front(self):
                return None

            async def evaluate(self, *_args, **_kwargs):
                self.evaluate_calls += 1

            async def close(self):
                return None

        page = Page()
        context = SimpleNamespace(pages=[page])

        rows, error = await fetch_douyin_search(
            SimpleNamespace(), SimpleNamespace(), "露营",
            max_results=1, captcha_wait_seconds=3, context=context,
        )

        self.assertEqual([row["aweme_id"] for row in rows], ["page-result-1"])
        self.assertEqual(error, "")
        self.assertEqual(page.evaluate_calls, 0)
        self.assertEqual(page.mouse.wheel.await_count, 0)

class KeywordCollectionPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "collection.db"
        db.init_db(str(self.db_path))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="douyin",
                nickname="collector",
                status="active",
                storage_state=json.dumps({
                    "cookies": [{
                        "name": "sessionid",
                        "value": "fixture-session",
                        "domain": ".douyin.com",
                        "path": "/",
                    }],
                    "origins": [],
                }),
            )
            session.add(account)
            session.commit()
            session.refresh(account)
            self.account = SimpleNamespace(
                id=account.id,
                platform=account.platform,
                storage_state=account.storage_state,
                proxy=account.proxy,
            )
            job = KeywordCollectionJob(
                platform="douyin",
                account_id=account.id,
                keywords=json.dumps(["露营"], ensure_ascii=False),
                max_contents_per_keyword=2,
                max_comments_per_content=2,
            )
            session.add(job)
            session.commit()
            session.refresh(job)
            self.job_id = job.id

    async def asyncTearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        self.temp_dir.cleanup()

    async def test_pipeline_persists_content_comments_and_deduplicates_retry(self):
        raw_aweme = {
            "aweme_id": "aweme-1",
            "desc": "周末露营清单",
            "create_time": 1_700_000_000,
            "author": {"nickname": "示例作者", "sec_uid": "author-1"},
            "statistics": {"digg_count": 88, "comment_count": 12},
            "video": {
                "play_addr": {"url_list": ["https://media.test/aweme-1.mp4"]},
                "cover": {"url_list": ["https://media.test/cover.jpg"]},
            },
        }
        comments = [
            {"comment_id": "c-1", "text": "想要链接", "user_nickname": "用户甲",
             "user_sec_uid": "MS4wLj-user-a", "ip_label": "北京",
             "like_count": 3, "create_time": 1_700_000_001, "reply_to": ""},
            {"comment_id": "c-2", "text": "已收藏", "user_nickname": "用户乙",
             "user_sec_uid": "MS4wLj-user-b", "ip_label": "上海",
             "like_count": 1, "create_time": 1_700_000_002, "reply_to": ""},
        ]
        headed_calls = []

        @asynccontextmanager
        async def temporary_headed_context(identity):
            headed_calls.append(identity)
            yield SimpleNamespace()

        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            temporary_headed_context=temporary_headed_context,
        )
        collector = KeywordCollector(Config(), browser, SimpleNamespace())
        collector._discover_douyin = AsyncMock(return_value=([raw_aweme], ""))
        collector._douyin_comments = AsyncMock(return_value=(comments, ""))

        first = await collector.run(self.job_id, self.account)
        second = await collector.run(self.job_id, self.account)

        self.assertEqual(first["contents"], 1)
        self.assertEqual(first["comments"], 2)
        self.assertEqual(second["contents"], 1)
        self.assertEqual(second["comments"], 2)
        with db.get_session() as session:
            contents = session.exec(select(KeywordCollectionContent)).all()
            saved_comments = session.exec(select(KeywordCollectionComment)).all()
            job = session.get(KeywordCollectionJob, self.job_id)
        self.assertEqual(len(contents), 1)
        self.assertEqual(len(saved_comments), 2)
        saved_by_id = {row.comment_id: row for row in saved_comments}
        self.assertEqual(saved_by_id["c-1"].user_sec_uid, "MS4wLj-user-a")
        self.assertEqual(saved_by_id["c-1"].ip_label, "北京")
        self.assertEqual(contents[0].author_id, "author-1")
        self.assertEqual(contents[0].collected_comment_count, 2)
        self.assertEqual(job.content_count, 1)
        self.assertEqual(job.comment_count, 2)
        self.assertEqual(headed_calls, [])

    async def test_browser_search_persists_verification_progress_and_checks_cancellation(self):
        context = SimpleNamespace()

        @asynccontextmanager
        async def headed(_identity):
            yield context

        cfg = Config()
        cfg.engine.douyin_read_mode = "browser"
        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            temporary_headed_context=headed,
        )
        collector = KeywordCollector(cfg, browser, SimpleNamespace())

        async def search(*_args, **kwargs):
            self.assertIs(kwargs["context"], context)
            self.assertFalse(kwargs["cancel_requested"]())
            kwargs["on_progress"]("等待人工完成抖音验证（剩余 300 秒）")
            with db.get_session() as session:
                job = session.get(KeywordCollectionJob, self.job_id)
                self.assertIn("等待人工", job.current_step)
                job.cancel_requested = True
                session.add(job)
                session.commit()
            self.assertTrue(kwargs["cancel_requested"]())
            return [], ""

        with patch("app.engine.collection.fetch_douyin_search", side_effect=search) as fetch:
            result = await collector.run(self.job_id, self.account)
        fetch.assert_awaited_once()
        self.assertEqual(result["errors"], 0)

    async def test_api_mode_does_not_create_a_headed_browser_context(self):
        raw_aweme = {
            "aweme_id": "api-aweme-1",
            "desc": "纯 API 搜索结果",
            "create_time": 1_700_000_000,
            "author": {"nickname": "示例作者", "sec_uid": "author-1"},
            "statistics": {"digg_count": 8, "comment_count": 0},
            "video": {"play_addr": {"url_list": ["https://media.test/api.mp4"]}},
        }
        headed_calls = []

        @asynccontextmanager
        async def temporary_headed_context(_identity):
            headed_calls.append(True)
            raise AssertionError("api mode must not create a browser context")
            yield  # pragma: no cover

        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            temporary_headed_context=temporary_headed_context,
        )
        cfg = Config()
        cfg.engine.douyin_read_mode = "api"
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        collector._discover_douyin = AsyncMock(return_value=([raw_aweme], ""))
        collector._douyin_comments = AsyncMock(return_value=([], ""))

        result = await collector.run(self.job_id, self.account)

        self.assertEqual(result["contents"], 1)
        self.assertEqual(headed_calls, [])
        self.assertIsNone(
            collector._discover_douyin.await_args.kwargs["context"])

    async def test_api_mode_comment_failure_never_falls_back_to_browser(self):
        cfg = Config()
        cfg.engine.douyin_read_mode = "api"
        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"))
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        client = SimpleNamespace(
            fetch_all_comments=AsyncMock(return_value=[]),
            last_error="empty_body",
        )
        aweme = Aweme(
            aweme_id="api-aweme-1", desc="", create_time=0,
            author_name="", media_type="video", comment_count=5)
        browser_fetch = AsyncMock(
            side_effect=AssertionError("api mode must not use browser comments"))

        with patch("app.engine.collection.fetch_comments", browser_fetch):
            comments, error = await collector._douyin_comments(
                self.account, client, aweme, 20, False)

        self.assertEqual(comments, [])
        self.assertIn("empty_body", error)
        client.fetch_all_comments.assert_awaited_once()
        browser_fetch.assert_not_awaited()

    async def test_hybrid_verification_required_never_opens_browser(self):
        identity = SimpleNamespace(ua="fixture-agent", account_id=self.account.id)
        fresh_calls = []

        @asynccontextmanager
        async def temporary_headed_context(value):
            fresh_calls.append(value)
            raise AssertionError("verification_required must not open a browser")
            yield

        browser = SimpleNamespace(
            identity_for=lambda _account: identity,
            temporary_headed_context=temporary_headed_context,
        )
        cfg = Config()
        cfg.engine.douyin_read_mode = "hybrid"
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        client = SimpleNamespace(
            search_awemes=AsyncMock(return_value=([], "verification_required")))
        job = SimpleNamespace(
            max_contents_per_keyword=3,
            max_pages_per_keyword=2,
            stagnant_pages=1,
            search_sort="general",
            publish_time="all",
            content_type="all",
            min_likes=0,
            min_comments=0,
            min_duration=0,
            max_duration=0,
        )
        browser_fetch = AsyncMock(side_effect=AssertionError(
            "verification_required must not fall back to browser search"))

        with patch("app.engine.collection.fetch_douyin_search", browser_fetch):
            values, error = await collector._discover_douyin(
                self.account, "电脑显卡", job, client)

        self.assertEqual(values, [])
        self.assertEqual(error, "抖音搜索接口失败: verification_required")
        self.assertEqual(fresh_calls, [])
        client.search_awemes.assert_awaited_once()
        browser_fetch.assert_not_awaited()

    async def test_hybrid_business_api_failure_can_fall_back_to_browser(self):
        context = SimpleNamespace()
        @asynccontextmanager
        async def temporary_headed_context(_identity):
            yield context

        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            temporary_headed_context=temporary_headed_context)
        cfg = Config()
        cfg.engine.douyin_read_mode = "hybrid"
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        client = SimpleNamespace(search_awemes=AsyncMock(
            return_value=([], "api_rejected:status_code=12345")))
        job = SimpleNamespace(
            max_contents_per_keyword=3, max_pages_per_keyword=2, stagnant_pages=1,
            search_sort="general", publish_time="all", content_type="all",
            min_likes=0, min_comments=0, min_duration=0, max_duration=0)
        browser_rows = [{"aweme_id": "browser-compat-result"}]
        browser_fetch = AsyncMock(return_value=(browser_rows, ""))

        with patch("app.engine.collection.fetch_douyin_search", browser_fetch):
            values, error = await collector._discover_douyin(
                self.account, "电脑显卡", job, client, context=context)

        self.assertEqual(values, browser_rows)
        self.assertEqual(error, "")
        client.search_awemes.assert_awaited_once()
        browser_fetch.assert_awaited_once()

    async def test_hybrid_compatibility_exception_can_fall_back_to_browser(self):
        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"))
        cfg = Config()
        cfg.engine.douyin_read_mode = "hybrid"
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        client = SimpleNamespace(search_awemes=AsyncMock(
            side_effect=ValueError("unexpected response shape")))
        job = SimpleNamespace(
            max_contents_per_keyword=3, max_pages_per_keyword=2, stagnant_pages=1,
            search_sort="general", publish_time="all", content_type="all",
            min_likes=0, min_comments=0, min_duration=0, max_duration=0)
        browser_rows = [{"aweme_id": "browser-exception-result"}]
        browser_fetch = AsyncMock(return_value=(browser_rows, ""))

        with patch("app.engine.collection.fetch_douyin_search", browser_fetch):
            values, error = await collector._discover_douyin(
                self.account, "电脑显卡", job, client)

        self.assertEqual(values, browser_rows)
        self.assertEqual(error, "")
        browser_fetch.assert_awaited_once()

    async def test_browser_mode_comments_never_call_direct_api(self):
        cfg = Config()
        cfg.engine.douyin_read_mode = "browser"
        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"))
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        client = SimpleNamespace(
            fetch_all_comments=AsyncMock(side_effect=AssertionError(
                "browser mode must not use direct comments API")),
            last_error="",
        )
        aweme = Aweme(
            aweme_id="browser-aweme-1", desc="", create_time=0,
            author_name="", media_type="video", comment_count=1)
        browser_fetch = AsyncMock(return_value=([{
            "cid": "comment-1",
            "text": "浏览器评论",
            "user": {"nickname": "用户甲", "uid": "user-1"},
            "create_time": 1_700_000_001,
        }], ""))

        with patch("app.engine.collection.fetch_comments", browser_fetch):
            comments, error = await collector._douyin_comments(
                self.account, client, aweme, 20, False)

        self.assertEqual(error, "")
        self.assertEqual(len(comments), 1)
        client.fetch_all_comments.assert_not_awaited()
        browser_fetch.assert_awaited_once()

    async def test_verification_required_stops_remaining_keywords_in_same_job(self):
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
            job.keywords = json.dumps(["露营", "穿搭"], ensure_ascii=False)
            session.add(job)
            session.commit()

        @asynccontextmanager
        async def temporary_headed_context(_identity):
            yield SimpleNamespace()

        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            temporary_headed_context=temporary_headed_context,
        )
        cfg = Config()
        cfg.engine.douyin_keyword_gap_seconds = 0
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        collector._discover_douyin = AsyncMock(return_value=(
            [], "抖音搜索接口失败: verification_required",
        ))

        result = await collector.run(self.job_id, self.account)

        self.assertEqual(collector._discover_douyin.await_count, 1)
        self.assertEqual(result["errors"], 1)
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
        self.assertIn("verification_required", job.error)


class KeywordCollectionEditTests(unittest.TestCase):
    def setUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "collection-edit.db"))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="douyin", nickname="编辑账号", status="active",
                storage_state='{"cookies":[{"name":"sessionid","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            job = KeywordCollectionJob(
                platform="douyin", account_id=account.id,
                keywords='["旧关键词"]', status="done", current_step="已完成",
            )
            session.add(job); session.commit(); session.refresh(job)
            content = KeywordCollectionContent(
                job_id=job.id, platform="douyin", keyword="旧关键词",
                aweme_id="keep-result",
            )
            session.add(content); session.commit()
            self.account_id, self.job_id = account.id, job.id

    def tearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    def _body(self, **changes):
        values = dict(
            platform="douyin", account_id=self.account_id,
            keywords=["新关键词", "第二个词"],
            max_contents_per_keyword=12,
            max_pages_per_keyword=24, stagnant_pages=4,
            search_sort="latest", publish_time="week", content_type="video",
            min_likes=100, min_comments=5,
            max_comments_per_content=30,
            include_replies=True, download_media=True,
            video_quality="1080", download_dir="",
        )
        values.update(changes)
        return KeywordCollectionIn(**values)

    def test_finished_task_can_be_edited_without_deleting_results(self):
        result = asyncio.run(update_keyword_collection(self.job_id, self._body()))

        self.assertEqual(result["keywords"], ["新关键词", "第二个词"])
        self.assertEqual(result["max_contents_per_keyword"], 12)
        self.assertEqual(result["max_pages_per_keyword"], 24)
        self.assertEqual(result["stagnant_pages"], 4)
        self.assertEqual(result["search_sort"], "latest")
        self.assertEqual(result["publish_time"], "week")
        self.assertEqual(result["content_type"], "video")
        self.assertEqual(result["min_likes"], 100)
        self.assertEqual(result["min_comments"], 5)
        self.assertEqual(result["max_comments_per_content"], 30)
        self.assertTrue(result["include_replies"])
        self.assertTrue(result["download_media"])
        self.assertEqual(result["video_quality"], "1080")
        self.assertEqual(result["status"], "done")
        self.assertEqual(result["current_step"], "配置已更新，可点击续跑")
        with db.get_session() as session:
            kept = session.exec(select(KeywordCollectionContent)).all()
        self.assertEqual([row.aweme_id for row in kept], ["keep-result"])

    def test_running_or_pending_task_must_be_canceled_before_edit(self):
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
            job.status = "running"
            session.add(job); session.commit()

        with self.assertRaises(HTTPException) as caught:
            asyncio.run(update_keyword_collection(self.job_id, self._body()))
        self.assertEqual(caught.exception.status_code, 409)

    def test_collection_depth_and_filter_values_are_validated(self):
        invalid_values = (
            {"max_pages_per_keyword": 0},
            {"stagnant_pages": 9},
            {"search_sort": "unknown"},
            {"publish_time": "month"},
            {"content_type": "live"},
            {"min_likes": -1},
        )
        for values in invalid_values:
            with self.subTest(values=values), self.assertRaises(HTTPException) as caught:
                asyncio.run(update_keyword_collection(self.job_id, self._body(**values)))
            self.assertEqual(caught.exception.status_code, 400)

    def test_xhs_collection_creation_uses_matching_logged_in_account(self):
        with db.get_session() as session:
            account = DouyinAccount(
                platform="xhs", nickname="小红书采集账号", status="active",
                storage_state='{"cookies":[{"name":"a1","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            xhs_account_id = account.id

        result = asyncio.run(create_keyword_collection(self._body(
            platform="xhs", account_id=xhs_account_id,
            search_sort="latest", content_type="images")))

        self.assertEqual(result["platform"], "xhs")
        self.assertEqual(result["account_id"], xhs_account_id)
        self.assertEqual(result["search_sort"], "latest")
        self.assertEqual(result["content_type"], "images")

    def test_kuaishou_collection_creation_uses_matching_logged_in_account(self):
        with db.get_session() as session:
            account = DouyinAccount(
                platform="kuaishou", nickname="快手采集账号", status="active",
                storage_state='{"cookies":[{"name":"passToken","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            account_id = account.id

        result = asyncio.run(create_keyword_collection(self._body(
            platform="kuaishou", account_id=account_id,
            search_sort="most_liked", content_type="video")))

        self.assertEqual(result["platform"], "kuaishou")
        self.assertEqual(result["account_id"], account_id)
        self.assertEqual(result["search_sort"], "most_liked")

    def test_channels_collection_only_accepts_current_account_works(self):
        with db.get_session() as session:
            account = DouyinAccount(
                platform="shipinhao", nickname="视频号采集账号", status="active",
                sec_uid="finder-fixture",
                storage_state='{"cookies":[{"name":"_finder_auth","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            account_id = account.id

        result = asyncio.run(create_keyword_collection(self._body(
            platform="shipinhao", account_id=account_id,
            source_type="account_works", source_name="", keywords=[],
            download_media=False)))

        self.assertEqual(result["platform"], "shipinhao")
        self.assertEqual(result["source_type"], "account_works")
        self.assertEqual(result["source_name"], "本账号作品")
        self.assertEqual(result["keywords"], [])
        filtered = asyncio.run(list_keyword_collections(platform="shipinhao"))
        self.assertEqual([job["id"] for job in filtered], [result["id"]])

        for source_type in ("keyword", "creator"):
            with self.subTest(source_type=source_type), self.assertRaises(HTTPException) as caught:
                asyncio.run(create_keyword_collection(self._body(
                    platform="shipinhao", account_id=account_id,
                    source_type=source_type, source_ref="someone")))
            self.assertEqual(caught.exception.status_code, 400)
            self.assertIn("本账号作品", caught.exception.detail)

    def test_finished_xhs_task_can_be_edited_but_not_moved_between_platforms(self):
        with db.get_session() as session:
            account = DouyinAccount(
                platform="xhs", nickname="小红书采集账号", status="active",
                storage_state='{"cookies":[{"name":"a1","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            job = KeywordCollectionJob(
                platform="xhs", account_id=account.id,
                keywords='["旧词"]', status="done")
            session.add(job); session.commit(); session.refresh(job)
            account_id, job_id = account.id, job.id

        updated = asyncio.run(update_keyword_collection(job_id, self._body(
            platform="xhs", account_id=account_id, search_sort="latest")))
        self.assertEqual(updated["platform"], "xhs")
        self.assertEqual(updated["search_sort"], "latest")

        with self.assertRaises(HTTPException) as caught:
            asyncio.run(update_keyword_collection(job_id, self._body()))
        self.assertEqual(caught.exception.status_code, 400)
        self.assertIn("不能变更平台", caught.exception.detail)


class XhsKeywordCollectionPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "xhs-collection.db"))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="xhs", nickname="collector", status="active",
                storage_state='{"cookies":[{"name":"a1","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            job = KeywordCollectionJob(
                platform="xhs", account_id=account.id,
                keywords='["露营"]', max_contents_per_keyword=2,
                max_pages_per_keyword=6, search_sort="latest",
                publish_time="week", content_type="images",
                min_likes=100, min_comments=2,
                max_comments_per_content=1,
            )
            session.add(job); session.commit(); session.refresh(job)
            self.account = SimpleNamespace(
                id=account.id, platform="xhs",
                storage_state=account.storage_state, proxy="")
            self.job_id = job.id

    async def asyncTearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    @staticmethod
    def _note(note_id: str, created: int, likes: int, comments: int) -> tuple[Aweme, dict, str]:
        aweme = Aweme(
            aweme_id=note_id, desc=note_id, create_time=created,
            author_name="作者", media_type="images",
            like_count=likes, comment_count=comments,
        )
        aweme.platform = "xhs"
        aweme.medias.append(MediaItem(
            url=f"https://media.test/{note_id}.jpg", kind="image",
            ext="jpg", index=0))
        return aweme, {
            "brief": {"note_id": note_id, "xsec_token": f"token-{note_id}",
                      "xsec_source": "web_search_result_note"},
            "card": {"note_id": note_id,
                     "user": {"user_id": f"author-{note_id}"}},
        }, ""

    async def test_filters_candidates_before_limit_and_preserves_access_context(self):
        now = int(__import__("time").time())
        raw = [{"id": value} for value in ("old", "low", "keep-1", "keep-2")]
        details = {
            "old": self._note("old", now - 30 * 86400, 900, 20),
            "low": self._note("low", now - 3600, 5, 20),
            "keep-1": self._note("keep-1", now - 1800, 500, 8),
            "keep-2": self._note("keep-2", now - 900, 300, 4),
        }
        browser = SimpleNamespace(
            identity_for=lambda _account: SimpleNamespace(ua="fixture-agent"),
            visible_page=lambda *_args, **_kwargs: None,
        )
        cfg = Config()
        cfg.engine.xhs_read_mode = "browser"
        cfg.engine.xhs_item_gap_seconds = 0
        collector = KeywordCollector(cfg, browser, SimpleNamespace())
        collector._discover_xhs_browser = AsyncMock(return_value=(raw, ""))
        collector._materialize_xhs_browser = AsyncMock(
            side_effect=lambda _identity, item: details[item["id"]])
        collector._xhs_comments_browser = AsyncMock(return_value=([{
            "comment_id": "comment-1", "text": "已收藏",
            "user_nickname": "用户", "like_count": 1,
            "create_time": now, "reply_to": "",
        }], ""))
        collector._xhs_gap = AsyncMock()

        result = await collector.run(self.job_id, self.account)

        self.assertEqual(result["contents"], 2)
        self.assertEqual(result["comments"], 2)
        self.assertEqual(collector._materialize_xhs_browser.await_count, 4)
        with db.get_session() as session:
            rows = session.exec(select(KeywordCollectionContent).order_by(
                KeywordCollectionContent.id)).all()
        self.assertEqual([row.aweme_id for row in rows], ["keep-1", "keep-2"])
        self.assertTrue(all(row.xsec_source == "web_search_result_note" for row in rows))

    async def test_api_search_forwards_xhs_sort_type_depth_and_stagnation(self):
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
            job.publish_time = "all"
            job.max_contents_per_keyword = 2
            job.content_type = "video"
            session.add(job); session.commit(); session.refresh(job)
        calls = []

        async def search_notes(keyword, **kwargs):
            calls.append((keyword, kwargs))
            return ([{"id": "b"}, {"id": "a"}] if kwargs["page"] == 1 else [])

        collector = KeywordCollector(Config(), SimpleNamespace(), SimpleNamespace())
        values, error = await collector._discover_xhs(
            SimpleNamespace(search_notes=search_notes), "露营", job)

        self.assertEqual(error, "")
        self.assertEqual(len(values), 2)
        self.assertEqual(calls[0][1]["sort"], "time_descending")
        self.assertEqual(calls[0][1]["note_type"], 1)

    def test_xhs_sort_and_filter_helpers_match_douyin_style_controls(self):
        rows = [
            {"id": "older", "note_card": {"time": 100,
             "interact_info": {"liked_count": "1.2万"}}},
            {"id": "newer", "note_card": {"time": 200,
             "interact_info": {"liked_count": "50"}}},
        ]
        self.assertEqual(
            [row["id"] for row in _sort_xhs_search_items(rows, "latest")],
            ["newer", "older"])
        self.assertEqual(
            [row["id"] for row in _sort_xhs_search_items(rows, "most_liked")],
            ["older", "newer"])
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
        accepted = self._note("accepted", int(__import__("time").time()), 100, 2)[0]
        rejected = self._note("rejected", int(__import__("time").time()), 99, 2)[0]
        self.assertTrue(_xhs_content_matches(accepted, job))
        self.assertFalse(_xhs_content_matches(rejected, job))


class KuaishouKeywordCollectionPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "ks-collection.db"))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="kuaishou", nickname="快手采集账号", status="active",
                storage_state='{"cookies":[{"name":"passToken","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            job = KeywordCollectionJob(
                platform="kuaishou", account_id=account.id,
                keywords='["露营"]', max_contents_per_keyword=1,
                max_pages_per_keyword=4, search_sort="most_liked",
                publish_time="week", content_type="video",
                min_likes=100, min_comments=2,
                max_comments_per_content=2, include_replies=True,
            )
            session.add(job); session.commit(); session.refresh(job)
            self.account = SimpleNamespace(
                id=account.id, platform="kuaishou",
                storage_state=account.storage_state, proxy="")
            self.job_id = job.id
        self.identity = SimpleNamespace(ua="fixture-agent")
        self.browser = SimpleNamespace(identity_for=lambda _account: self.identity)
        self.collector = KeywordCollector(
            Config(), self.browser, SimpleNamespace())

    async def asyncTearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    @staticmethod
    def _feed(photo_id: str, likes: int, comments: int, *, age: int = 60):
        return {
            "photo": {
                "id": photo_id, "caption": photo_id,
                "timestamp": (int(__import__("time").time()) - age) * 1000,
                "realLikeCount": likes, "commentCount": comments,
                "photoUrl": f"https://media.test/{photo_id}.mp4",
            },
            "author": {"id": f"author-{photo_id}", "name": "快手作者"},
        }

    async def test_keyword_pipeline_filters_persists_and_collects_replies(self):
        self.collector._discover_ks = AsyncMock(return_value=([
            self._feed("low", 5, 20),
            self._feed("keep", 500, 8),
        ], ""))
        self.collector._ks_comments = AsyncMock(return_value=([{
            "comment_id": "root", "text": "根评论", "user_nickname": "甲",
            "like_count": 2, "create_time": 100, "reply_to": "",
        }, {
            "comment_id": "reply", "text": "回复", "user_nickname": "乙",
            "like_count": 1, "create_time": 101, "reply_to": "root",
        }], ""))

        with patch("app.engine.collection.asyncio.sleep", new=AsyncMock()):
            result = await self.collector.run(self.job_id, self.account)

        self.assertEqual(result["contents"], 1)
        self.assertEqual(result["comments"], 2)
        self.collector._discover_ks.assert_awaited_once()
        self.collector._ks_comments.assert_awaited_once_with(
            self.identity, "keep", 2, True)
        with db.get_session() as session:
            row = session.exec(select(KeywordCollectionContent)).one()
        self.assertEqual(row.aweme_id, "keep")
        self.assertEqual(row.author_id, "author-keep")
        self.assertEqual(row.platform, "kuaishou")
        self.assertEqual(
            _collection_content_dict(row)["url"],
            "https://www.kuaishou.com/short-video/keep")

    async def test_creator_source_uses_profile_feed_path(self):
        with db.get_session() as session:
            job = session.get(KeywordCollectionJob, self.job_id)
            job.source_type = "creator"
            job.source_ref = "3xCREATOR"
            job.source_name = "创作者 3xCREATOR"
            job.max_comments_per_content = 0
            job.min_likes = 0
            job.min_comments = 0
            job.publish_time = "all"
            session.add(job); session.commit()
        self.collector._discover_ks = AsyncMock(
            side_effect=AssertionError("creator source must not use search"))
        self.collector._discover_ks_creator = AsyncMock(return_value=([
            self._feed("creator-work", 20, 1),
        ], ""))

        with patch("app.engine.collection.asyncio.sleep", new=AsyncMock()):
            result = await self.collector.run(self.job_id, self.account)

        self.assertEqual(result["contents"], 1)
        self.collector._discover_ks_creator.assert_awaited_once()
        self.collector._discover_ks.assert_not_awaited()

    def test_kuaishou_sort_uses_timestamp_and_like_count(self):
        old_popular = self._feed("old", 900, 1, age=3600)
        new_quiet = self._feed("new", 10, 1, age=10)
        self.assertEqual(
            [item["photo"]["id"] for item in _sort_ks_search_items(
                [old_popular, new_quiet], "latest")], ["new", "old"])
        self.assertEqual(
            [item["photo"]["id"] for item in _sort_ks_search_items(
                [new_quiet, old_popular], "most_liked")], ["old", "new"])


class ChannelsAccountCollectionPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "channels-collection.db"))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="shipinhao", nickname="视频号作者", sec_uid="finder-owner",
                status="active",
                storage_state='{"cookies":[{"name":"_finder_auth","value":"fixture"}]}',
            )
            session.add(account); session.commit(); session.refresh(account)
            job = KeywordCollectionJob(
                platform="shipinhao", account_id=account.id,
                source_type="account_works", source_name="本账号作品", keywords="[]",
                max_contents_per_keyword=2, max_pages_per_keyword=4,
                content_type="all", min_likes=100, min_comments=1,
                max_comments_per_content=3, include_replies=True, download_media=True,
            )
            session.add(job); session.commit(); session.refresh(job)
            self.account = SimpleNamespace(
                id=account.id, platform="shipinhao", nickname=account.nickname,
                sec_uid=account.sec_uid, storage_state=account.storage_state, proxy="")
            self.job_id = job.id
        self.identity = SimpleNamespace(ua="fixture-agent")
        self.browser = SimpleNamespace(identity_for=lambda _account: self.identity)
        self.collector = KeywordCollector(Config(), self.browser, SimpleNamespace())

    async def asyncTearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    @staticmethod
    def _work(object_id: str, likes: int, comments: int, media_type: int):
        return {
            "objectId": object_id, "createTime": int(__import__("time").time()),
            "likeCount": likes, "commentCount": comments,
            "desc": {
                "description": object_id, "mediaType": media_type,
                "media": [{"mediaType": media_type, "thumbUrl": f"https://media.test/{object_id}.jpg"}],
            },
        }

    async def test_account_works_persist_video_photo_comments_and_metadata_only_media(self):
        self.collector._discover_channels_account = AsyncMock(return_value=([
            self._work("below-threshold", 10, 3, 2),
            self._work("video-work", 500, 2, 1),
            self._work("photo-work", 300, 1, 2),
        ], ""))
        self.collector._channels_comments = AsyncMock(side_effect=[([
            {"comment_id": "root", "text": "根评论", "user_nickname": "甲",
             "user_sec_uid": "finder-a", "like_count": 2, "create_time": 100,
             "reply_to": ""},
            {"comment_id": "reply", "text": "回复", "user_nickname": "乙",
             "user_sec_uid": "finder-b", "like_count": 1, "create_time": 101,
             "reply_to": "root"},
        ], ""), ([], "")])

        with patch("app.engine.collection.asyncio.sleep", new=AsyncMock()):
            result = await self.collector.run(self.job_id, self.account)

        self.assertEqual(result["contents"], 2)
        self.assertEqual(result["comments"], 2)
        self.collector._discover_channels_account.assert_awaited_once()
        self.assertIs(
            self.collector._discover_channels_account.await_args.args[0], self.identity)
        self.assertEqual(
            self.collector._discover_channels_account.await_args.args[1].id, self.job_id)
        self.assertEqual(self.collector._channels_comments.await_count, 2)
        with db.get_session() as session:
            rows = session.exec(select(KeywordCollectionContent).order_by(
                KeywordCollectionContent.id)).all()
            comments = session.exec(select(KeywordCollectionComment).order_by(
                KeywordCollectionComment.id)).all()
        self.assertEqual([row.aweme_id for row in rows], ["video-work", "photo-work"])
        self.assertEqual([row.media_type for row in rows], ["video", "images"])
        self.assertTrue(all(row.download_status == "skipped" for row in rows))
        self.assertTrue(all("仅保存作品元数据" in row.error for row in rows))
        self.assertTrue(all(row.author_name == "视频号作者" for row in rows))
        self.assertTrue(all(row.author_id == "finder-owner" for row in rows))
        self.assertEqual({row.user_sec_uid for row in comments}, {"finder-a", "finder-b"})
        self.assertEqual(
            _collection_content_dict(rows[0])["url"],
            "https://channels.weixin.qq.com/platform/post/list")


class KeywordCollectionMediaPreviewTests(unittest.TestCase):
    def setUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "collection-preview.db"))
        with db.get_session() as session:
            job = KeywordCollectionJob(
                platform="douyin", account_id=1, keywords='["预览"]', status="done",
            )
            session.add(job); session.commit(); session.refresh(job)
            self.job_id = job.id

    def tearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    def _add_content(self, **values):
        defaults = dict(
            job_id=self.job_id, platform="douyin", keyword="预览",
            aweme_id="preview-1", desc="可预览视频", media_type="video",
            media_json=json.dumps([{
                "kind": "video", "ext": "mp4",
                "url": "https://media.test/preview-1.mp4",
            }]),
            download_status="done",
        )
        defaults.update(values)
        with db.get_session() as session:
            content = KeywordCollectionContent(**defaults)
            session.add(content); session.commit(); session.refresh(content)
            return content.id

    def test_downloaded_video_is_exposed_to_result_list_and_preview(self):
        media = Path(self.temp_dir.name) / "preview-1_title.mp4"
        media.write_bytes(b"fixture-video")
        content_id = self._add_content(local_path=str(media))

        with db.get_session() as session:
            item = _collection_content_dict(session.get(KeywordCollectionContent, content_id))
        preview = asyncio.run(keyword_collection_content_media(self.job_id, content_id))
        response = asyncio.run(keyword_collection_local_media(self.job_id, content_id, 0))

        self.assertTrue(item["local_exists"])
        self.assertTrue(item["preview_available"])
        self.assertEqual(item["media_count"], 1)
        self.assertEqual(item["file_size"], len(b"fixture-video"))
        self.assertEqual(
            preview["local_url"],
            f"/api/collections/{self.job_id}/contents/{content_id}/local-media/0",
        )
        self.assertEqual(preview["medias"][0]["kind"], "video")
        self.assertEqual(Path(response.path), media.resolve())

    def test_gallery_preview_only_includes_files_for_current_aweme(self):
        folder = Path(self.temp_dir.name) / "author"
        folder.mkdir()
        (folder / "gallery-1_title_1.jpg").write_bytes(b"second")
        (folder / "gallery-1_title_0.jpg").write_bytes(b"first")
        (folder / "another-work_title_0.jpg").write_bytes(b"other")
        content_id = self._add_content(
            aweme_id="gallery-1", media_type="images", local_path=str(folder),
            media_json="[]",
        )

        preview = asyncio.run(keyword_collection_content_media(self.job_id, content_id))

        self.assertEqual(len(preview["medias"]), 2)
        self.assertTrue(all(item["kind"] == "image" for item in preview["medias"]))
        self.assertTrue(preview["medias"][0]["url"].endswith("/local-media/0"))
        self.assertEqual(preview["local_url"], "")

    def test_preview_rejects_content_from_another_job(self):
        content_id = self._add_content()
        with self.assertRaises(HTTPException) as caught:
            asyncio.run(keyword_collection_content_media(self.job_id + 99, content_id))
        self.assertEqual(caught.exception.status_code, 404)


class KeywordCollectionCommentQueryTests(unittest.TestCase):
    def setUp(self):
        self.previous_engine = db._engine
        self.temp_dir = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.temp_dir.name) / "collection-comments.db"))
        with db.get_session() as session:
            job = KeywordCollectionJob(
                platform="douyin", account_id=1, keywords='["露营"]', status="done",
            )
            session.add(job); session.commit(); session.refresh(job)
            self.job_id = job.id
            self.contents = []
            for keyword, aweme_id, desc in (("露营", "work-1", "海边露营"), ("骑行", "work-2", "城市骑行")):
                content = KeywordCollectionContent(
                    job_id=self.job_id, platform="douyin", keyword=keyword,
                    aweme_id=aweme_id, desc=desc,
                )
                session.add(content); session.commit(); session.refresh(content)
                self.contents.append(content.id)
            comments = [
                (self.contents[0], "u-1", "线索评论", "广东", 100),
                (self.contents[1], "u-1", "另一条", "北京", 200),
                (self.contents[1], "u-2", "普通评论", "广东", 150),
                (self.contents[0], "", "无稳定 ID", "广东", 300),
            ]
            for content_id, uid, text, region, created in comments:
                session.add(KeywordCollectionComment(
                    job_id=self.job_id, content_id=content_id, platform="douyin",
                    aweme_id="work-1" if content_id == self.contents[0] else "work-2",
                    comment_id=f"comment-{created}", user_sec_uid=uid, text=text,
                    ip_label=region, create_time=created,
                ))
            session.commit()

    def tearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.temp_dir.cleanup()

    def test_filters_pagination_and_cross_work_user_dedupe(self):
        filtered = asyncio.run(list_keyword_collection_comments(
            self.job_id, q="评论", ip_label="广东", paginate=True, page_size=1))
        self.assertEqual(filtered["total"], 2)
        self.assertEqual(filtered["pages"], 2)
        self.assertIn(filtered["items"][0]["source_keyword"], {"露营", "骑行"})

        deduped = asyncio.run(list_keyword_collection_comments(
            self.job_id, dedupe_users=True, paginate=True, page_size=20))
        self.assertEqual(deduped["total"], 2)
        self.assertEqual({row["user_sec_uid"] for row in deduped["items"]}, {"u-1", "u-2"})
        self.assertEqual(next(row for row in deduped["items"] if row["user_sec_uid"] == "u-1")["create_time"], 200)

    def test_invalid_content_scope_and_filtered_export(self):
        with self.assertRaises(HTTPException) as caught:
            asyncio.run(list_keyword_collection_comments(self.job_id, content_id=99999))
        self.assertEqual(caught.exception.status_code, 404)

        response = asyncio.run(export_keyword_collection(self.job_id, q="线索"))

        async def read_response():
            return b"".join([chunk async for chunk in response.body_iterator])

        workbook = load_workbook(BytesIO(asyncio.run(read_response())))
        self.assertEqual(workbook["作品"].max_row, 2)
        self.assertEqual(workbook["评论"].max_row, 2)
        self.assertEqual(workbook["用户线索"]["A2"].value, "u-1")


class KeywordCollectionReportTests(unittest.TestCase):
    def test_report_has_overview_content_and_comment_sheets(self):
        job = KeywordCollectionJob(
            id=7,
            platform="xhs",
            account_id=1,
            keywords='["露营"]',
            status="done",
        )
        content = KeywordCollectionContent(
            id=8,
            job_id=7,
            platform="xhs",
            keyword="露营",
            aweme_id="note-1",
            desc="=FORMULA",
            collected_comment_count=1,
        )
        comment = KeywordCollectionComment(
            id=9,
            job_id=7,
            content_id=8,
            platform="xhs",
            aweme_id="note-1",
            comment_id="comment-1",
            text="+comment",
            user_sec_uid="MS4wLj-report-user",
            ip_label="广东",
        )

        workbook = load_workbook(BytesIO(
            build_keyword_collection_report(job, [content], [comment])
        ))

        self.assertEqual(workbook.sheetnames, ["概览", "作品", "评论", "用户线索"])
        self.assertEqual(workbook["作品"]["E2"].value, "'=FORMULA")
        self.assertEqual(workbook["评论"]["F2"].value, "'+comment")
        self.assertEqual(workbook["评论"]["H2"].value, "MS4wLj-report-user")
        self.assertEqual(workbook["评论"]["I2"].value, "广东")
        self.assertEqual(workbook["用户线索"]["A2"].value, "MS4wLj-report-user")
        self.assertEqual(workbook["用户线索"]["D2"].value, 1)


if __name__ == "__main__":
    unittest.main()
